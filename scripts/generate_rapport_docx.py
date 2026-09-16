#!/usr/bin/env python3
"""Convertit docs/rapport-projet.md en docs/rapport-projet.docx (python-docx).

Sous-ensemble markdown géré : titres h1-h3, paragraphes, gras/italique/code inline,
listes à puces et numérotées, tables pipe, blocs de code clôturés, images,
séparateurs horizontaux. Les photos des membres sont intégrées.
Usage : python3 scripts/generate_rapport_docx.py
"""
import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MD = os.path.join(ROOT, "docs", "rapport-projet.md")
OUT = os.path.join(ROOT, "docs", "rapport-projet.docx")

BLUE = RGBColor(0x0D, 0x3B, 0x66)
BLUE2 = RGBColor(0x14, 0x54, 0x8C)


def add_runs(par, text):
    """Ajoute les runs inline (gras, italique, code) d'une ligne markdown."""
    token_re = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")
    for tok in token_re.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            r = par.add_run(tok[2:-2])
            r.bold = True
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            r = par.add_run(tok[1:-1])
            r.italic = True
        elif tok.startswith("`") and tok.endswith("`"):
            r = par.add_run(tok[1:-1])
            r.font.name = "Courier New"
            r.font.size = Pt(9)
        else:
            par.add_run(tok)


def add_table(doc, rows):
    """Ajoute une table pipe markdown (avec ligne d'en-tête séparée par ---)."""
    parsed = []
    for row in rows:
        cells = [c.strip() for c in row.strip().strip("|").split("|")]
        parsed.append(cells)
    parsed = [r for r in parsed if not all(re.fullmatch(r":?-{2,}:?", c or "---") for c in r)]
    if not parsed:
        return
    ncols = max(len(r) for r in parsed)
    table = doc.add_table(rows=len(parsed), cols=ncols)
    table.style = "Light Grid Accent 1"
    for i, row in enumerate(parsed):
        for j in range(ncols):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            txt = row[j] if j < len(row) else ""
            add_runs(p, txt)
            for r in p.runs:
                r.font.size = Pt(9)
                if i == 0:
                    r.bold = True


def main():
    with open(MD, "r", encoding="utf-8") as f:
        lines = f.read().split("\n")

    doc = Document()

    # Styles de base
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    doc.styles["Heading 1"].font.color.rgb = BLUE
    doc.styles["Heading 2"].font.color.rgb = BLUE
    doc.styles["Heading 3"].font.color.rgb = BLUE2

    i = 0
    while i < len(lines):
        line = lines[i]

        # Bloc de code clôturé
        if line.strip().startswith("```"):
            i += 1
            code_lines = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            r = p.add_run("\n".join(code_lines))
            r.font.name = "Courier New"
            r.font.size = Pt(7.5)
            i += 1
            continue

        # Table pipe
        if line.strip().startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|[\s:|-]+\|?\s*$", lines[i + 1] or ""):
            rows = [line]
            i += 1
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i])
                i += 1
            add_table(doc, rows)
            doc.add_paragraph()
            continue

        # Image markdown
        m = re.match(r"^!\[(.*)\]\((.+)\)\s*$", line.strip())
        if m:
            alt, rel = m.groups()
            img = os.path.join(ROOT, "docs", rel)
            if os.path.exists(img):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(img, height=Cm(4.2))
            i += 1
            continue

        # Titres
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            level = min(len(m.group(1)), 4)
            text = re.sub(r"\*\*|\*", "", m.group(2))
            h = doc.add_heading(text, level=level)
            for r in h.runs:
                r.font.color.rgb = BLUE if level <= 2 else BLUE2
            i += 1
            continue

        # Séparateur
        if re.match(r"^---+\s*$", line):
            i += 1
            continue

        # Liste à puces / numérotée
        m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", line)
        if m:
            indent, marker, content = m.groups()
            style = "List Bullet" if marker in "-*" else "List Number"
            p = doc.add_paragraph(style=style)
            p.paragraph_format.left_indent = Cm(1.0 + 0.5 * (len(indent) // 2))
            add_runs(p, content)
            i += 1
            continue

        # Citation
        if line.strip().startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            add_runs(p, line.strip().lstrip("> "))
            for r in p.runs:
                r.italic = True
            i += 1
            continue

        # Ligne vide
        if not line.strip():
            i += 1
            continue

        # Paragraphe simple
        p = doc.add_paragraph()
        add_runs(p, line.strip())
        i += 1

    doc.save(OUT)
    print(f"OK -> {OUT} ({os.path.getsize(OUT)/1024:.0f} Ko)")


if __name__ == "__main__":
    main()
